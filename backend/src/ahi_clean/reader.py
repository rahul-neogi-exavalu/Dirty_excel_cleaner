"""Read a workbook into a plain grid of Python values, one grid per sheet.

Everything downstream works on lists of lists rather than openpyxl objects, so
the structural heuristics stay easy to test with hand-written fixtures.
"""

from __future__ import annotations

import os
import re
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.reader.excel import ExcelReader
from openpyxl.styles.stylesheet import apply_stylesheet
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple

from . import delimited


@dataclass
class SheetGrid:
    """A sheet's cells as a dense rectangular grid, plus what we noticed reading it."""

    name: str
    rows: list[list]
    styles: list[list] = field(default_factory=list)
    error_cells: list[str] = field(default_factory=list)
    formula_cells: list[str] = field(default_factory=list)
    uncached_formula_cells: list[str] = field(default_factory=list)
    merged_ranges: list[str] = field(default_factory=list)
    image_count: int = 0
    source_format: str = "xlsx"
    # Flags about how the file itself was read (separator, encoding); every column of
    # every table from this grid carries them. See flags.py.
    read_flags: list[str] = field(default_factory=list)

    @property
    def height(self) -> int:
        return len(self.rows)

    @property
    def width(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def error_lines(self, transposed: bool = False) -> dict[int, list[str]]:
        """0-based row (or column, for a sheet read sideways) -> its Excel error values.

        Error cells are nulled on read so they cannot anchor a table, which leaves a row
        holding only errors indistinguishable from a blank one. It is not blank -- it held
        something -- so it has to be accounted for as its own kind of removal.

        A merged error cell covers every line of its range: File2's ``#VALUE!`` sits in
        A1:D2, which Excel shows centred across rows 1-2, so neither row is blank.
        """
        spans = []
        for merged in self.merged_ranges:
            try:
                spans.append(range_boundaries(merged))
            except ValueError:
                continue
        lines: dict[int, list[str]] = {}
        for entry in self.error_cells:
            reference, _, value = entry.partition("=")
            try:
                row, column = coordinate_to_tuple(reference)
            except ValueError:
                continue
            first = last = column if transposed else row
            for min_col, min_row, max_col, max_row in spans:
                if (min_row, min_col) == (row, column):
                    first, last = (min_col, max_col) if transposed else (min_row, max_row)
                    break
            for line in range(first, last + 1):
                lines.setdefault(line - 1, []).append(value)
        return lines


# A sheet beyond this many cells is refused rather than attempted. A documented
# ceiling is a feature; an unexplained hang that swaps the machine is not.
DEFAULT_MAX_CELLS = 20_000_000


class SheetTooLarge(Exception):
    """A sheet exceeds the configured cell ceiling."""


DELIMITED_SUFFIXES = {".csv", ".tsv", ".txt"}


def sheet_names(path) -> list[str]:
    """The worksheets of ``path`` in workbook order, without reading any cells."""
    if Path(path).suffix.lower() in DELIMITED_SUFFIXES:
        return [Path(path).stem]
    with open_workbook(path) as workbook:
        return [worksheet.title for worksheet in workbook.worksheets]


def read_workbook(path, max_cells: int = DEFAULT_MAX_CELLS, only=None) -> list[SheetGrid]:
    """Load the sheets of ``path`` into :class:`SheetGrid` s -- every sheet, or ``only`` these.

    A delimited file is one sheet by definition and takes a different reader, but yields
    the same grid -- which is the point of the grid being a list of lists. Nothing
    downstream needs to know where the cells came from.
    """
    if Path(path).suffix.lower() in DELIMITED_SUFFIXES:
        return delimited.read_delimited(path, max_cells, SheetTooLarge)
    with _Source(path) as source:
        return [source.sheet(name, max_cells) for name in source.names if only is None or name in only]


def read_sheet(path, name: str, max_cells: int = DEFAULT_MAX_CELLS, reuse: bool = False) -> SheetGrid:
    """Load one sheet, touching no other. Reading sheets independently is what lets them
    be read side by side, each in its own process.

    ``reuse`` keeps the file's parsed structure (shared strings, styles) for the next
    sheet of the same file -- for a worker process, which reads one sheet at a time from
    a single thread. Not for use from several threads at once.
    """
    if Path(path).suffix.lower() in DELIMITED_SUFFIXES:
        return delimited.read_delimited(path, max_cells, SheetTooLarge)[0]
    with _Source(path, reuse=reuse) as source:
        return source.sheet(name, max_cells)


class _LazySheet(ReadOnlyWorksheet):
    """A read-only worksheet that does not measure itself when opened.

    openpyxl sizes each read-only sheet from its ``<dimension>`` tag, and when a writer
    leaves that tag out it parses the *whole sheet* to find out -- for every sheet, on
    every open. Opening one sheet of an eight-sheet file cost as much as reading all
    eight. The reader streams real rows and measures the used extent itself, so the
    probe is pure waste.
    """

    def _get_size(self):
        pass


@contextmanager
def open_workbook(path):
    """Open a workbook read-only, values not formulas, without touching any sheet's cells."""
    workbook, archive = _open_structure(path)
    try:
        yield workbook
    finally:
        archive.close()


def _open_structure(path):
    """openpyxl's read-only load, minus the per-sheet size probe (see ``_LazySheet``)."""
    # data_only=True gives cached formula results rather than formula text.
    reader = ExcelReader(path, read_only=True, data_only=True, keep_links=False)
    try:
        reader.read_manifest()
        reader.read_strings()
        reader.read_workbook()
        reader.read_properties()
        reader.read_custom()
        reader.read_theme()
        apply_stylesheet(reader.archive, reader.wb)
        for sheet, rel in reader.parser.find_sheets():
            if rel.target not in reader.valid_files or "chartsheet" in rel.Type:
                continue
            worksheet = _LazySheet(reader.wb, sheet.name, rel.target, reader.shared_strings)
            worksheet.sheet_state = sheet.state
            reader.wb._sheets.append(worksheet)
    except BaseException:
        reader.archive.close()
        raise
    return reader.wb, reader.archive


# One parsed workbook structure per worker process, for the next sheet of the same file.
_reused: dict = {}


def _file_key(path) -> tuple:
    stat = os.stat(path)
    return (str(Path(path).resolve()), stat.st_size, stat.st_mtime_ns)


class _Source:
    """An open workbook, streamed a sheet at a time.

    openpyxl's full mode parses *every* sheet into cell objects before returning, and
    this module used to load the file twice that way (values, then formula text). The
    parse is the cost -- about four seconds per 150,000 cells -- so it is now paid once,
    and only for the sheets asked for: values stream through read-only mode, and the
    two things read-only mode cannot see, formula cells and merged ranges, come from one
    fast scan of the same sheet's XML.
    """

    def __init__(self, path, reuse: bool = False):
        self.path = path
        key = _file_key(path) if reuse else None
        workbook = _reused.get(key) if reuse else None
        if workbook is None:
            workbook, self.archive = _open_structure(path)
            if reuse:
                _reused.clear()
                _reused[key] = workbook
        else:
            # The structure is kept, never the file handle: an open handle would stop
            # the upload being deleted on Windows.
            self.archive = zipfile.ZipFile(path)
            workbook._archive = self.archive
        self.workbook = workbook
        self.names = [worksheet.title for worksheet in workbook.worksheets]

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.archive.close()

    def sheet(self, name: str, max_cells: int) -> SheetGrid:
        worksheet = self.workbook[name]
        formulas, merged, positional = _scan_sheet_xml(self.archive, worksheet._worksheet_path)
        grid = _read_sheet(worksheet, merged, max_cells)
        if not positional:
            # A writer may leave out cell references; then only openpyxl knows where a
            # formula sits, so ask it -- slowly, and only for such a sheet.
            formulas = _formula_cells_by_openpyxl(self.path, name)
        _note_formulas(grid, formulas)
        return grid


# One cell element and its reference; the body (group 2) is absent for <c .../>.
_CELL = re.compile(rb"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
_REF = re.compile(rb'\br="([A-Z]{1,3}[0-9]+)"')
_MERGE = re.compile(rb'<mergeCell\b[^>]*?\bref="([A-Z0-9:]+)"')
_SCAN_CHUNK = 4 * 1024 * 1024


def _scan_sheet_xml(archive: zipfile.ZipFile, part: str) -> tuple[list[str], list[str], bool]:
    """Formula cell references and merged ranges, from the raw sheet XML.

    Streamed in chunks so a very large sheet is never held whole. A chunk is cut just
    before its last ``<c `` so every cell scanned is complete; merged ranges live after
    the cell data and are complete by the time they are reached. Text is XML-escaped, so
    a literal ``<f`` inside a cell can only be a formula element. The third value is False
    if any cell carried no reference, meaning positions cannot be trusted from here.
    """
    formulas: list[str] = []
    merged: list[str] = []
    positional = True
    tail = b""

    def scan(block: bytes) -> None:
        nonlocal positional
        for match in _CELL.finditer(block):
            body = match.group(2)
            if body is None or b"<f" not in body:
                continue
            ref = _REF.search(match.group(1))
            if ref is None:
                positional = False
            else:
                formulas.append(ref.group(1).decode())
        merged.extend(ref.decode() for ref in _MERGE.findall(block))

    with archive.open(part.lstrip("/")) as handle:
        while chunk := handle.read(_SCAN_CHUNK):
            block = tail + chunk
            cut = block.rfind(b"<c ")
            if cut <= 0:
                tail = block
                continue
            scan(block[:cut])
            tail = block[cut:]
    scan(tail)
    return formulas, merged, positional


def _formula_cells_by_openpyxl(path, name: str) -> list[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        return [
            cell.coordinate
            for row in workbook[name].iter_rows()
            for cell in row
            if getattr(cell, "data_type", None) == "f"
        ]
    finally:
        workbook.close()


def _note_formulas(grid: SheetGrid, coordinates: list[str]) -> None:
    """Record formula cells, and which of them have no saved result.

    data_only gives the value Excel cached the last time it saved; a workbook written by
    a library and never opened in Excel has no cache, so a formula-driven column arrives
    silently empty. Knowing which cells hold formulas is the difference between
    reporting that and shipping a column of nulls as if the source were blank.
    """
    for coordinate in coordinates:
        grid.formula_cells.append(coordinate)
        column, row = coordinate_to_tuple(coordinate)[::-1]
        inside = row <= grid.height and column <= len(grid.rows[row - 1])
        if not inside or grid.rows[row - 1][column - 1] is None:
            grid.uncached_formula_cells.append(coordinate)


def count_embedded_images(path) -> int:
    """Count images in the drawing layer.

    Embedded logos are anchored to drawings, not to cells, so they are invisible
    to any value-based scan -- File2's banner occupies rows that read as empty.
    They never affect extraction; this is recorded for the audit log so a reviewer
    can see why the top of that sheet looked blank.
    """
    # A delimited file is not a zip and cannot carry a drawing layer at all.
    if Path(path).suffix.lower() in DELIMITED_SUFFIXES:
        return 0
    with zipfile.ZipFile(path) as archive:
        return sum(1 for name in archive.namelist() if name.startswith("xl/media/"))


def _read_sheet(worksheet, merged_ranges: list[str], max_cells: int = DEFAULT_MAX_CELLS) -> SheetGrid:
    """Stream one read-only worksheet into a dense grid.

    Only the genuinely used rectangle is kept, not the one the file declares: stray
    formatting inflates the declared size freely -- a single styled cell at row 100,000
    makes a ten-row report claim a hundred thousand rows. Rows are gathered ragged, the
    extent is taken from the last non-empty value, and only then is the grid made dense.
    """
    # The declared <dimension> is a hint writers get wrong; stream the real rows.
    worksheet.reset_dimensions()
    raw_rows: list[list] = []
    raw_styles: list[list] = []
    error_cells: list[str] = []
    style_cache: dict = {}
    height = width = 0

    for row_index, cells in enumerate(worksheet.iter_rows(), start=1):
        values: list = []
        styles: list = []
        for column_index, cell in enumerate(cells, start=1):
            value = cell.value
            if value is not None:
                # Extent follows the raw value, before trimming or nulling below, so a
                # cell holding only spaces or an error still counts as used.
                height = row_index
                width = max(width, column_index)
            # Excel error cells ('#VALUE!', '#REF!', ...) come back as ordinary
            # strings. Left alone they read as populated and defeat blank-row
            # detection -- File2!A1 sits under the logo and would anchor the
            # header hunt to row 1. Null them here, once, before anything scores.
            if cell.data_type == "e":
                error_cells.append(f"{get_column_letter(column_index)}{row_index}={value}")
                value = None
            if isinstance(value, str):
                value = value.strip() or None
            values.append(value)
            styles.append(None if value is None else _style_of(cell, style_cache))
        raw_rows.append(values)
        raw_styles.append(styles)
        if height * width > max_cells:
            raise SheetTooLarge(
                f"sheet '{worksheet.title}' spans at least {height} x {width} = {height * width:,} "
                f"cells, above the {max_cells:,} ceiling"
            )

    del raw_rows[height:], raw_styles[height:]
    rows = [_fit(values, width) for values in raw_rows]
    styles = [_fit(values, width) for values in raw_styles]
    _propagate_merged_values(merged_ranges, rows)

    return SheetGrid(
        name=worksheet.title,
        rows=rows,
        styles=styles,
        error_cells=error_cells,
        merged_ranges=list(merged_ranges),
    )


def _fit(values: list, width: int) -> list:
    return values[:width] + [None] * (width - len(values)) if len(values) != width else values


def _style_of(cell, cache: dict) -> dict:
    """Visual emphasis on a cell, used as one header-detection signal.

    Report writers style header rows and leave data rows plain, so this is free
    evidence already sitting in the file. It is only ever additive: a workbook with no
    formatting scores zero here and is decided by the content signals alone.

    A workbook has a handful of distinct cell styles and every cell names one by index,
    so the answer is worked out once per style and shared. Resolving font, fill and
    border objects afresh for every cell was most of the time it took to read a sheet.
    The shared dicts are never mutated downstream.
    """
    key = getattr(cell, "_style_id", None)
    if key is None:
        return _emphasis(cell)
    found = cache.get(key)
    if found is None:
        found = cache[key] = _emphasis(cell)
    return found


def _emphasis(cell) -> dict:
    font = cell.font
    fill = cell.fill
    filled = bool(
        fill is not None
        and fill.fill_type not in (None, "none")
        and getattr(fill.fgColor, "rgb", None) not in (None, "00000000")
    )
    bordered = bool(cell.border is not None and cell.border.bottom is not None
                    and cell.border.bottom.style)
    return {
        "bold": bool(font is not None and font.bold),
        "filled": filled,
        "bordered": bordered,
        "emphasised": bool((font is not None and font.bold) or filled or bordered),
    }


def _propagate_merged_values(merged_ranges: list[str], rows: list[list]) -> None:
    """Fill a merged range with its top-left value so the grid stays rectangular.

    A range reaching past the used extent (a merged banner wider than the table) is
    filled only as far as the grid goes.
    """
    height = len(rows)
    for merged in merged_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(merged)
        if min_row > height or min_col > len(rows[min_row - 1]):
            continue
        top_left = rows[min_row - 1][min_col - 1]
        if top_left is None:
            continue
        for row_index in range(min_row, min(max_row, height) + 1):
            row = rows[row_index - 1]
            for column_index in range(min_col, min(max_col, len(row)) + 1):
                row[column_index - 1] = top_left
