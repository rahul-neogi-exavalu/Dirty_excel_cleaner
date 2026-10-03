"""Write the 14 test workbooks for File -> Bronze -> Silver, plus test LOTL rows.

    python backend/tools/make_silver_test_files.py

Deterministic (fixed seed), so the files and their expected outcomes never drift. Each
file drives one scenario of the two AHI scenario documents; enhancements/test-files/
README.md says which, and what the app should do with it. Ingest them in number order.
"""

from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[2] / "enhancements" / "test-files"

BASE = [
    "Profit Center Name", "Profit Center Number", "Insurance Company Name", "Producer Name",
    "Policy Number", "Premium", "Accounting Effective Date", "Policy Effective Date",
    "Transaction Effective Date",
]
# pc_id -> (legacy office name, profit center number): the LOTL seed.
OFFICES = {
    "0094": ("Dayton Office", "0094"),
    "0101": ("Cleveland Office", "0101"),
    "0202": ("Akron Office", "0202"),
    "0303": ("Columbus Office", "0303"),
    "0404": ("Cincinnati Office", "0404"),
    "0505": ("Canton Office", "0505"),
    "0515": ("Toledo Office", "0515"),
    "0606": ("Youngstown Office", "0606"),
    "0707": ("Lima Office", "0707"),
    "0909": ("Mansfield Office", "0909"),
}
CARRIERS = ["Heritage Casualty Co", "National Indemnity Co", "Global Assurance Corp", "Liberty Mutual", "Zenith Insurance Group"]
PRODUCERS = ["Pinnacle Agency Partners", "Apex Insurance Brokers", "Metro Agency Group", "Coastal Risk Advisors", "Brown & Brown"]

rng = random.Random(20260701)


def record(pc_id: str, month: int, n: int, year: int = 2026) -> dict:
    name, number = OFFICES[pc_id]
    accounting = date(year, month, 1 + (n * 3) % 27)
    return {
        "Profit Center Name": name,
        "Profit Center Number": int(number),
        "Insurance Company Name": rng.choice(CARRIERS),
        "Producer Name": rng.choice(PRODUCERS),
        "Policy Number": f"POL-{pc_id}-{month:02d}{n:03d}",
        "Premium": round(rng.uniform(800, 25000), 2),
        "Accounting Effective Date": accounting,
        "Policy Effective Date": accounting - timedelta(days=rng.randint(10, 300)),
        "Transaction Effective Date": accounting - timedelta(days=rng.randint(0, 5)),
    }


def rows(pc_id: str, months, per_month: int, start: int = 0) -> list[dict]:
    return [record(pc_id, m, start + i) for m in months for i in range(per_month)]


def sheet(book, name: str, header: list[str], data: list[list], hidden: bool = False):
    ws = book.create_sheet(name)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for values in data:
        ws.append(values)
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 22
    if hidden:
        ws.sheet_state = "hidden"
    return ws


def save(book, filename: str) -> None:
    if "Sheet" in book.sheetnames and len(book.sheetnames) > 1:
        del book["Sheet"]
    book.save(OUT / filename)
    print("wrote", filename)


def simple(filename: str, sheet_name: str, header: list[str], data: list[dict], rename=None) -> None:
    rename = rename or {}
    book = openpyxl.Workbook()
    sheet(book, sheet_name, [rename.get(h, h) for h in header], [[r.get(h) for h in header] for r in data])
    save(book, filename)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    # 1. Base load, Jan-Jun.
    simple("01_ARR_pc0101_2026_JanJun.xlsx", "ARR", BASE, rows("0101", range(1, 7), 5))
    # 2. July, identical columns.
    simple("02_ARR_pc0101_2026_Jul.xlsx", "ARR", BASE, rows("0101", [7], 5, start=100))
    # 3. August, same columns in another order.
    reordered = [BASE[i] for i in (4, 0, 1, 5, 2, 3, 8, 6, 7)]
    simple("03_ARR_pc0101_2026_Aug_reordered.xlsx", "ARR", reordered, rows("0101", [8], 5, start=200))

    # 4. September: renamed + extra columns (drives every cascade stage).
    september = rows("0101", [9], 5, start=300)
    for r in september:
        r["Notes"] = rng.choice(["", "Renewal", "Endorsement", "New business"])
    header = BASE + ["Notes"]
    simple("04_ARR_pc0101_2026_Sep_newcols.xlsx", "ARR", header, september,
           rename={"Premium": "Premium Amt", "Insurance Company Name": "Carrier",
                   "Producer Name": "Writing Agency"})

    # 5. Second source: month-named sheets, its own header wording.
    wording = {
        "Profit Center Name": "PC Name", "Profit Center Number": "PC No",
        "Insurance Company Name": "Carrier Name", "Producer Name": "Agency",
        "Policy Number": "Policy #", "Premium": "Written Premium",
        "Accounting Effective Date": "Acct Eff Date", "Policy Effective Date": "Pol Eff Date",
        "Transaction Effective Date": "Trans Eff Date",
    }
    book = openpyxl.Workbook()
    for month, label in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"], start=1):
        data = rows("0202", [month], 4, start=month * 10)
        sheet(book, label, [wording[h] for h in BASE], [[r[h] for h in BASE] for r in data])
    save(book, "05_Prem_pc0202_2026.xlsx")

    # 6. The four profit-center cases from the Bronze -> Silver doc, plus a no-match row.
    cases = rows("0303", [1, 2, 3], 2)
    plan = [
        ("Dayton Office", 94),        # both present, consistent      -> kept, 0094
        ("Toledo Office", None),      # number missing                -> filled_number 0515
        ("Dayton Office", 515),       # name/number mismatch          -> corrected 0094
        (None, None),                 # both missing                  -> filled_name via pc_id 0303
        ("Unknown Office", 777),      # not in LOTL                   -> no_match, 0777
        ("Toledo Office", 515),       # both present, consistent      -> kept
    ]
    for r, (name, number) in zip(cases, plan):
        r["Profit Center Name"], r["Profit Center Number"] = name, number
    simple("06_PC_pc0303_2026_cases.xlsx", "Cases", BASE, cases)

    # 7. Every date format in the doc, Excel serials and invalid dates, as text.
    formats = [
        "2026-01-15 10:30:00", "2026-01-16", "01/17/2026", "1/8/2026", "20260119",
        "01-20-2026", "1-9-2026", 46030, "not a date", "13/45/2026",
    ]
    dated = rows("0404", [1], len(formats))
    for r, value in zip(dated, formats):
        r["Accounting Effective Date"] = value
    simple("07_Dates_pc0404_2026.xlsx", "Dates", BASE, dated)

    # 8. Decimal and string formats: currency, separators, accounting negatives, junk,
    #    padded text and blank cells.
    amounts = ["1200.50", "1,200.50", "$1,200.50", "(250.00)", "abc", " 3,400.10 ", "$ 99.00", "-75.25"]
    money = rows("0505", [2], len(amounts))
    for i, (r, value) in enumerate(zip(money, amounts)):
        r["Premium"] = value
        r["Producer Name"] = "" if i % 3 == 0 else f"  {r['Producer Name']}  "
    simple("08_Money_pc0505_2026.xlsx", "Money", BASE, money)

    # 9. Revised Jan-Jun for pc0101: overlaps file 1 only -> Bronze REPLACE of file 1.
    simple("09_ARR_pc0101_2026_JanJun_revised.xlsx", "ARR", BASE, rows("0101", range(1, 7), 4, start=500))

    # 10. A dirty report: titles, subtotals, blank rows, a hidden sheet and a second
    #     small table beside the main one.
    book = openpyxl.Workbook()
    ws = book.create_sheet("Report")
    ws.append(["Youngstown Office - Premium Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated 2026-08-01"])
    ws.append([])
    ws.append(BASE + [None, None, "Producer Code", "Producer Name"])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    lookup = [(f"P{100 + i}", name) for i, name in enumerate(PRODUCERS)]
    line = 0
    for month in range(1, 8):
        block = rows("0606", [month], 3, start=month * 10)
        for r in block:
            extra = list(lookup[line]) if line < len(lookup) else [None, None]
            ws.append([r[h] for h in BASE] + [None, None] + extra)
            line += 1
        ws.append([None, None, None, None, f"Subtotal {month:02d}", round(sum(r["Premium"] for r in block), 2)])
        ws.append([])
    sheet(book, "Notes", ["Note"], [["Internal - do not load"]], hidden=True)
    save(book, "10_Report_pc0606_2026_JanJul.xlsx")

    # 11. May-Jul for pc0101: overlaps the loaded Jan-Jun file only partly. Neither
    #     Replace (loses Jan-Apr) nor Append (doubles May-Jun) is safe: the reviewer decides.
    simple("11_ARR_pc0101_2026_MayJul_overlap.xlsx", "ARR", BASE, rows("0101", [5, 6, 7], 2, start=700))

    # 12. October: the columns in another order *and* one more (cases 2 and 3 at once):
    #     schema evolution, not a separate table.
    october = rows("0101", [10], 4, start=800)
    for r in october:
        r["Commission %"] = f"{rng.randint(5, 15)}%"
    shuffled = [BASE[i] for i in (8, 7, 6, 5, 4, 3, 2, 1, 0)] + ["Commission %"]
    simple("12_ARR_pc0101_2026_Oct_reorder_newcol.xlsx", "ARR", shuffled, october)

    # 13. Money and dates the document's formats do not cover: a two-digit year, fractional
    #     seconds, ISO 'T', an out-of-range serial, currency with parentheses, half-cent
    #     rounding, a European number, a code with letters.
    edge = rows("0909", [3], 6)
    # Premium (kept as text by the cleaner, so Silver reads it): -250.00, 1.01, 0.13, NULL,
    # NULL, 9999999999999999.99. The dates are typed by the cleaner before Silver sees them.
    for r, (premium, accounting) in zip(edge, [
        ("$(250.00)", "1/5/26"),
        ("1.005", "2026-03-04 13:45:00.123"),
        ("0.125", "2026-03-05T08:00:00"),
        ("1 200,50", "99999"),
        ("USD 100", "45000"),
        ("9999999999999999.99", "03/06/2026"),
    ]):
        r["Premium"], r["Accounting Effective Date"] = premium, accounting
    simple("13_MoneyDates_pc0909_2026.xlsx", "Edge", BASE, edge)

    # 14. Two headers that clean to the same name (Amount ($) -> amount, Amount (%) ->
    #     amount_2), then February lists them the other way round: each must still land
    #     in its own column.
    clash = ["Policy Number", "Amount ($)", "Amount (%)", "Accounting Effective Date"]
    january = [[f"POL-0707-01{n:03d}", 100 + n, 0.5, date(2026, 1, 10 + n)] for n in range(4)]
    february = [[f"POL-0707-02{n:03d}", 0.5, 200 + n, date(2026, 2, 10 + n)] for n in range(4)]
    book = openpyxl.Workbook()
    sheet(book, "Clash", clash, january)
    save(book, "14_Clash_pc0707_2026_Jan.xlsx")
    book = openpyxl.Workbook()
    sheet(book, "Clash", [clash[0], clash[2], clash[1], clash[3]], february)
    save(book, "14_Clash_pc0707_2026_Feb_swapped.xlsx")

    # Test offices in the LOTL's own shape, added on top of the business's LOTL with
    # backend/tools/seed_reference.py --add-lotl (the real LOTL does not know them).
    with open(OUT / "lotl_seed.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["profit_center_number", "legacy_office_name", "status"])
        for number, (name, _) in OFFICES.items():
            writer.writerow([number.lstrip("0"), name, "Active"])
    print("wrote lotl_seed.csv")


if __name__ == "__main__":
    main()
