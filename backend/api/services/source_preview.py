"""The uploaded file as it is, before cleaning: one sheet's cells, paged by rows and columns.

The cleaner reads a sheet very differently -- it trims text, nulls error cells, fills
merged ranges, finds the table and drops banners, subtotals and blank rows. A preview
shows none of that. Cells keep their sheet positions (row numbers and column letters as
Excel shows them), so the reviewer can see exactly what cleaning is about to change.

A sheet is read once and held, so paging through it does not re-read the file. What is
held is bounded: a sheet past ``SOURCE_PREVIEW_MAX_CELLS`` keeps only its first rows for
paging (the rest are counted, not shown), and the sheets held together share
``SOURCE_PREVIEW_CACHE_CELLS``, the least recently viewed dropped first.
"""

from __future__ import annotations

import csv
import datetime as _dt
import math
import threading
from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl.utils import get_column_letter, range_boundaries

from ahi_clean import delimited, failures, reader

from .. import config
from ..errors import ApiError, not_found
from ..store import Upload


@dataclass
class RawSheet:
    """One sheet's values, trimmed to the last row and column that hold anything."""

    rows: list[tuple]  # the rows held for paging, each without its trailing empty cells
    height: int  # used rows: the sheet row of the last value
    width: int  # used columns
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # min_col, min_row, max_col, max_row
    source_format: str = "xlsx"

    @property
    def cells(self) -> int:
        return sum(len(row) for row in self.rows)


_cache: OrderedDict[tuple[str, str], RawSheet] = OrderedDict()
_cache_lock = threading.Lock()
# One lock per sheet being read, so a burst of page requests reads the file once.
_loading: dict[tuple[str, str], threading.Lock] = {}


def preview(upload: Upload, sheet: str, offset: int, limit: int, col_offset: int, col_limit: int) -> dict:
    info = next((item for item in upload.sheets if item["name"] == sheet), None)
    if info is None:
        raise not_found(f"A sheet named '{sheet}' in {upload.filename}")
    raw = _sheet(upload, sheet)
    columns = range(col_offset, min(col_offset + col_limit, raw.width))
    window = raw.rows[offset: offset + limit]
    rows = [[_json_value(row[index]) if index < len(row) else None for index in columns] for row in window]
    first_row, last_row = offset + 1, offset + len(window)
    first_col, last_col = col_offset + 1, col_offset + len(columns)
    return {
        "sheet": sheet,
        "hidden": bool(info.get("hidden")),
        "source_format": raw.source_format,
        "total_rows": raw.height,
        "total_columns": raw.width,
        # Fewer than total_rows when the sheet is too large to hold whole.
        "available_rows": len(raw.rows),
        "offset": offset,
        "limit": limit,
        "col_offset": col_offset,
        "col_limit": col_limit,
        "columns": [get_column_letter(index + 1) for index in columns],
        "rows": rows,
        "merged": [
            _reference(*bounds) for bounds in raw.merged
            if bounds[1] <= last_row and bounds[3] >= first_row and bounds[0] <= last_col and bounds[2] >= first_col
        ],
        "merged_total": len(raw.merged),
    }


def forget(upload_id: str) -> None:
    """Drop a removed upload's sheets."""
    with _cache_lock:
        for key in [key for key in _cache if key[0] == upload_id]:
            del _cache[key]


def _sheet(upload: Upload, name: str) -> RawSheet:
    key = (upload.id, name)
    with _cache_lock:
        found = _cache.get(key)
        if found is not None:
            _cache.move_to_end(key)
            return found
        lock = _loading.setdefault(key, threading.Lock())
    with lock:
        try:
            with _cache_lock:
                found = _cache.get(key)
            if found is None:
                found = _read(upload, name)
                with _cache_lock:
                    _cache[key] = found
                    _cache.move_to_end(key)
                    _evict()
        finally:
            with _cache_lock:
                _loading.pop(key, None)
    return found


def _evict() -> None:
    """Keep the held sheets within budget; the one just read always stays."""
    total = sum(sheet.cells for sheet in _cache.values())
    while total > config.SOURCE_PREVIEW_CACHE_CELLS and len(_cache) > 1:
        _, dropped = _cache.popitem(last=False)
        total -= dropped.cells


def _read(upload: Upload, name: str) -> RawSheet:
    try:
        if upload.kind == "delimited":
            return _read_delimited(upload)
        with reader.raw_sheet(upload.path, name) as (rows, merged):
            sheet = _collect(rows)
        sheet.merged = [range_boundaries(ref) for ref in merged]
        return sheet
    except KeyError as error:
        raise not_found(f"A sheet named '{name}' in {upload.filename}") from error
    except FileNotFoundError as error:
        raise ApiError(404, "not_found", f"{upload.filename} is no longer on the server.",
                       "Upload the file again.") from error
    except ApiError:
        raise
    except Exception as error:  # noqa: BLE001 - classified the way the cleaner would
        failure = failures.classify(upload.path, error)
        raise ApiError(422, failure.kind, f"{upload.filename} could not be previewed.",
                       failure.advice, failure.detail) from error


def _read_delimited(upload: Upload) -> RawSheet:
    text, encoding = delimited.read_text(upload.path)
    delimiter = delimited.sniff_delimiter(text[: delimited.SNIFF_BYTES])
    rows = csv.reader(text.splitlines(), delimiter=delimiter)
    sheet = _collect(tuple(value if value != "" else None for value in row) for row in rows)
    sheet.source_format = (f"{_DELIMITERS.get(delimiter, repr(delimiter))}-separated text, "
                           f"{_ENCODINGS.get(encoding, encoding)}")
    return sheet


_DELIMITERS = {",": "Comma", ";": "Semicolon", "\t": "Tab", "|": "Pipe"}
_ENCODINGS = {"utf-8-sig": "UTF-8", "utf-8": "UTF-8", "cp1252": "Windows-1252", "latin-1": "Latin-1"}


def _collect(rows) -> RawSheet:
    """Hold the rows up to the cell ceiling; keep counting the used extent after it."""
    held: list[tuple] = []
    cells = height = width = 0
    full = False
    for number, values in enumerate(rows, start=1):
        used = _used_length(values)
        if used:
            height, width = number, max(width, used)
        if not full:
            held.append(tuple(values[:used]))
            cells += used
            full = cells >= config.SOURCE_PREVIEW_MAX_CELLS
    del held[height:]  # trailing empty rows
    return RawSheet(rows=held, height=height, width=width)


def _used_length(values) -> int:
    """Cells up to and including the last one holding something."""
    for index in range(len(values) - 1, -1, -1):
        if values[index] is not None and values[index] != "":
            return index + 1
    return 0


def _reference(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _json_value(value):
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, _dt.datetime):
        # A date cell reads back as midnight; show it as the date Excel shows.
        return value.date().isoformat() if value.time() == _dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, (_dt.timedelta, Decimal)):
        return str(value)
    return str(value)
