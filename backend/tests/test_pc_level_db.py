"""A profit center's files at its own level, end to end on Postgres: a broker list that came
with the transactions, joined with them in Silver; and a summary the profit center sent
already aggregated, loaded as reported into silver_aggregate.

Skipped unless AHI_TEST_DATABASE_URL points at a database the test may write to. Each run
uses its own throwaway schemas and drops them afterwards.
"""

import io
import os
import sys
import time
import uuid
from datetime import date
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

URL = os.environ.get("AHI_TEST_DATABASE_URL")
pytestmark = pytest.mark.skipif(not URL, reason="AHI_TEST_DATABASE_URL is not set")

psycopg = pytest.importorskip("psycopg")
from fastapi.testclient import TestClient  # noqa: E402
from sse_client import validate  # noqa: E402

# The business's required columns, but MarketProvider: the broker list carries it.
HEADERS = ["AccountingEffective Date", "CommissionPct", "GrossCommissionAmount", "InsuranceCompany Name",
           "PolicyEffectiveDate", "PolicyNumber", "Premium", "Producer/Agency Name", "ProducerCommissionAmount",
           "ProducerCommissionPct", "Revenue", "Transaction Detail", "TransactionEffectiveDate"]
BROKERS = [("Agency 0", "Market A", "East"), ("Agency 1", "Market B", "West"), ("Agency 2", "Market A", "North"),
           ("Agency 3", "Market C", "South"), ("Agency 9", "Market B", "East")]


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    from psycopg.conninfo import conninfo_to_dict

    from api import config, db

    parts = conninfo_to_dict(URL)
    suffix = uuid.uuid4().hex[:8]
    work = tmp_path_factory.mktemp("workspace")
    names = ("WORK_DIR", "UPLOAD_DIR", "JOB_DIR", "DB_HOST", "DB_PORT", "DB_NAME", "DB_USER", "DB_PASSWORD",
             "DB_SSLMODE", "BRONZE_SCHEMA", "CONTROL_SCHEMA", "SILVER_SCHEMA", "CLEANSED_SCHEMA", "STAGING_SCHEMA",
             "LOTL_TABLE", "DB_CONFIGURED", "INGEST_ENABLED", "AI_ENABLED", "WORD2VEC_PATH")
    saved = {name: getattr(config, name) for name in names}
    config.WORK_DIR, config.UPLOAD_DIR, config.JOB_DIR = work, work / "uploads", work / "jobs"
    config.DB_HOST, config.DB_PORT = parts.get("host", "localhost"), int(parts.get("port", 5432))
    config.DB_NAME, config.DB_USER = parts.get("dbname", "postgres"), parts.get("user", "postgres")
    config.DB_PASSWORD, config.DB_SSLMODE = parts.get("password", ""), parts.get("sslmode", "disable")
    config.BRONZE_SCHEMA, config.CONTROL_SCHEMA = f"bronze_p{suffix}", f"ingest_p{suffix}"
    config.SILVER_SCHEMA, config.CLEANSED_SCHEMA = f"silver_p{suffix}", f"cleansed_p{suffix}"
    config.STAGING_SCHEMA, config.LOTL_TABLE = f"staging_p{suffix}", ""
    config.DB_CONFIGURED = config.INGEST_ENABLED = True
    config.AI_ENABLED, config.WORD2VEC_PATH = False, ""  # deterministic: no external calls
    db.close()
    from api.main import app

    yield TestClient(app), config
    db.close()
    with psycopg.connect(URL, autocommit=True) as conn:
        for schema in (config.BRONZE_SCHEMA, config.CONTROL_SCHEMA, config.SILVER_SCHEMA, config.CLEANSED_SCHEMA,
                       config.STAGING_SCHEMA):
            conn.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
    for name, value in saved.items():
        setattr(config, name, value)


# --- helpers -------------------------------------------------------------------------------


def _save(book) -> bytes:
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _transactions() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(HEADERS)
    n = 0
    for month in range(1, 7):
        for index in range(4):
            n += 1
            values = {"AccountingEffective Date": date(2026, month, 1 + n % 27), "CommissionPct": 0.12,
                      "GrossCommissionAmount": 120.5 + n, "InsuranceCompany Name": f"Carrier {n % 5}",
                      "PolicyEffectiveDate": date(2026, month, 1), "PolicyNumber": f"POL-{month:02d}{n:04d}",
                      "Premium": 1000.5 + n, "Producer/Agency Name": f"agency {n % 4}",  # case differs from the list
                      "ProducerCommissionAmount": 60 + n, "ProducerCommissionPct": 0.06, "Revenue": 60.5,
                      "Transaction Detail": "New Business", "TransactionEffectiveDate": date(2026, month, 1 + n % 27)}
            sheet.append([values[name] for name in HEADERS])
    return _save(book)


def _broker_list() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Brokers"
    sheet.append(["Account_Name", "MarketProvider", "Region", "Agents"])
    for index, row in enumerate(BROKERS, 1):
        sheet.append(list(row) + [index * 3])
    return _save(book)


def _summary() -> bytes:
    """A premium summary a month to a sheet, spaced for print like PC2030's."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    partners = [("Accident Fund", 9928.42, 5593.17, 443.81), ("CompSource Mutual", 2643.97, 1375.23, 207.54),
                ("Convex UK", 0, 4036.57, 0), ("Hastings", 6538.9, 0, 925.17), ("Tokio Marine", 0, 3640.43, 242.84)]
    for month in ("January 2026", "February 2026", "March 2026"):
        sheet = book.create_sheet(month)
        if month.startswith("February"):
            sheet.append([None, "NO PREMIUM WRITTEN IN FEBRUARY"])
            continue
        for row in (["Premium Summary by Partner"], [month], [], [],
                    ["Delegated Authority Partner", "Casualty Treaty", None, "Property Treaty", None, "Workers Comp",
                     None, "TOTALS"], [], []):
            sheet.append(row)
        for name, casualty, prop, workers in partners:
            sheet.append([name, casualty, None, prop, None, workers, None, round(casualty + prop + workers, 2)])
            sheet.append([])
        sums = [round(sum(item[index] for item in partners), 2) for index in (1, 2, 3)]
        sheet.append(["TOTALS", sums[0], None, sums[1], None, sums[2], None, round(sum(sums), 2)])
    return _save(book)


def _wait(client, url, active=("queued", "running", "draft")):
    for _ in range(600):
        state = client.get(url).json()
        if state["status"] not in active:
            return state
        time.sleep(0.1)
    raise AssertionError(f"{url} did not finish")


def _clean(client, name: str, content: bytes) -> str:
    upload = client.post("/api/workbooks", files={"file": (name, content)}).json()
    job = client.post("/api/jobs", json={"workbook_id": upload["id"], "sheets": [s["name"] for s in upload["sheets"]]})
    job = job.json()
    assert _wait(client, f"/api/jobs/{job['id']}", ("queued", "running"))["status"] == "succeeded"
    return job["id"]


def _patch(client, session: dict, key: str, change: dict) -> dict:
    response = client.patch(f"/api/validations/{session['id']}/outputs/{key}", json=change)
    assert response.status_code == 200, response.text
    return response.json()


def _ingest(client) -> dict:
    plan = client.post("/api/bronze/plans", json={})
    assert plan.status_code == 201, plan.text
    plan = plan.json()
    keys = [item["key"] for item in plan["items"]]
    assert client.post(f"/api/bronze/plans/{plan['id']}/approve", json={"confirmed": keys}).status_code == 202
    state = _wait(client, f"/api/bronze/plans/{plan['id']}", ("draft", "running"))
    assert state["status"] == "succeeded", state
    return plan


def _q(config, statement: str, *args):
    with psycopg.connect(URL) as conn:
        return conn.execute(statement.format(b=config.BRONZE_SCHEMA, c=config.CONTROL_SCHEMA, s=config.SILVER_SCHEMA,
                                             st=config.STAGING_SCHEMA), args).fetchall()


def _approve(client, run: dict) -> dict:
    response = client.post(f"/api/silver/runs/{run['id']}/approve", json={})
    assert response.status_code == 202, response.text
    state = _wait(client, f"/api/silver/runs/{run['id']}")
    assert state["status"] == "succeeded", state.get("error")
    return state


# --- a broker list that came with the transactions -------------------------------------------


def test_a_companion_list_is_joined_with_the_transactions_in_silver(env):
    client, config = env
    session = validate(client, [_clean(client, "PC0001_Transactional_08082026.xlsx", _transactions()),
                                _clean(client, "PC0001_BrokerList_08082026.xlsx", _broker_list())])
    by_file = {file["job_id"]: file["file_name"] for file in session["files"]}
    outputs = {by_file[o["job_id"]].split("_")[1]: o for o in session["outputs"]}
    transactions, brokers = outputs["Transactional"], outputs["BrokerList"]
    assert transactions["missing"] == ["MarketProvider"] and brokers["options"] == ["companion"]

    # The broker list came with the transactions: checked together, it is fit.
    session = _patch(client, session, brokers["key"], {"choice": "companion", "companion_of": transactions["key"]})
    outputs = {o["key"]: o for o in session["outputs"]}
    assert outputs[transactions["key"]]["verdict"] == "ready" and outputs[brokers["key"]]["action"] == "APPEND"
    staged = client.post(f"/api/validations/{session['id']}/stage", json={"keys": [brokers["key"]]})
    assert staged.status_code == 200, staged.text
    rows = {row[1]: row for row in _q(config, "SELECT control_id, file_name, companion_of, processing_action, "
                                              "staging_table FROM {c}.control_table WHERE file_name LIKE 'PC0001_%%'")}
    data_row, list_row = rows["PC0001_Transactional_08082026.xlsx"], rows["PC0001_BrokerList_08082026.xlsx"]
    assert list_row[2] == data_row[0] and data_row[2] is None
    assert (data_row[4], list_row[4]) == ("ext_pc0001_data_stg", "ext_pc0001_brokers_stg")

    # Ingest: two tables; the list never replaces the transactions.
    plan = _ingest(client)
    assert sorted((item["table_name"], item["action"]) for item in plan["items"]) == [
        ("ext_pc0001_brokers", "create"), ("ext_pc0001_data", "create")]

    # Silver: the transactions alone are selected; the list comes along to be joined.
    loads = client.get("/api/silver/eligible").json()["loads"]
    data_load = next(load["ingestion_id"] for load in loads if load["table_name"] == "ext_pc0001_data")
    run = client.post("/api/silver/runs", json={"ingestion_ids": [data_load]}).json()
    assert any("came with a selected file" in note for note in run["notes"])
    link = run["links"][0]
    assert (link["left_table"], link["right_table"]) == ("ext_pc0001_data", "ext_pc0001_brokers")
    best = link["keys"][0]
    assert (best["left"], best["right"]) == ("producer_agency_name", "account_name") and "overlap" in best["methods"]
    assert any("join the two" in blocker for blocker in run["blockers"])

    response = client.patch(f"/api/silver/runs/{run['id']}/join", json={
        "left_table": "ext_pc0001_data", "right_table": "ext_pc0001_brokers", "how": "left",
        "keys": [{"left": "producer_agency_name", "right": "account_name"}]})
    assert response.status_code == 200, response.text
    run = response.json()
    tables = {table["table_name"]: table for table in run["tables"]}
    data = tables["ext_pc0001_data"]
    assert tables["ext_pc0001_brokers"]["joined_into"] == "ext_pc0001_data"
    assert data["join"]["stats"]["matched"] == 24 and data["join"]["stats"]["unmatched"] == 0  # case aside
    assert "ext_pc0001_brokers.marketprovider" in data["columns"]
    # The list's MarketProvider is voted onto market_provider by its own name.
    mapped = {row["bronze_column"]: row["silver_column"] for row in data["mapping"]}
    assert mapped["ext_pc0001_brokers.marketprovider"] == "market_provider"
    assert not run["blockers"], run["blockers"]

    _approve(client, run)
    providers = _q(config, "SELECT DISTINCT market_provider FROM {s}.silver_transaction WHERE source_system = 'pc0001'")
    assert {row[0] for row in providers} == {"Market A", "Market B", "Market C"}
    assert _q(config, "SELECT count(*) FROM {s}.silver_transaction WHERE source_system = 'pc0001'")[0][0] == 24
    statuses = _q(config, "SELECT table_name, silver_status FROM {c}.ingestion WHERE table_name LIKE 'ext_pc0001_%%'")
    assert {row[1] for row in statuses} == {"succeeded"}
    assert _q(config, "SELECT how, keys FROM {c}.silver_join") == [
        ("left", [{"left": "Producer/Agency Name", "right": "Account_Name"}])]


# --- a summary the profit center aggregated -------------------------------------------------------


def test_an_aggregated_summary_goes_to_the_aggregate_table_as_reported(env):
    client, config = env
    session = validate(client, [_clean(client, "PC0002_Premium Summary January 1 2026 to March 31 2026_04152026.xlsx",
                                       _summary())])
    output = session["outputs"][0]
    assert output["aggregated"] and output["verdict"] == "ready" and output["origin"] == "sheets"
    staged = client.post(f"/api/validations/{session['id']}/stage", json={})
    assert staged.status_code == 200, staged.text
    control = _q(config, "SELECT is_aggregated, staging_table, date_detail FROM {c}.control_table "
                         "WHERE file_name LIKE 'PC0002_%%'")
    assert control == [("Y", "ext_pc0002_data_agg_stg", None)]
    assert _q(config, "SELECT count(*) FROM {st}.ext_pc0002_data_agg_stg")[0][0] == 12  # 2 sheets x 6 rows

    plan = _ingest(client)
    assert [item["table_name"] for item in plan["items"]] == ["ext_pc0002_data_agg"]
    months = _q(config, "SELECT _reporting_month, count(*) FROM {b}.ext_pc0002_data_agg GROUP BY 1 ORDER BY 1")
    assert months == [("2026-01", 6), ("2026-03", 6)]
    assert _q(config, "SELECT is_aggregated FROM {c}.ingestion WHERE table_name = 'ext_pc0002_data_agg'") == [("Y",)]

    loads = client.get("/api/silver/eligible").json()["loads"]
    load = next(item for item in loads if item["table_name"] == "ext_pc0002_data_agg")
    assert load["aggregated"]
    run = client.post("/api/silver/runs", json={"ingestion_ids": [load["ingestion_id"]]}).json()
    table = run["tables"][0]
    assert table["aggregated"] and table["rollups"] == ["totals"]
    assert table["spread"]["columns"] == ["casualty_treaty", "property_treaty", "workers_comp"]
    assert table["spread"]["measure"] == "premium"  # the "Premium Summary"
    mapped = {row["bronze_column"]: row["silver_column"] for row in table["mapping"]}
    assert mapped == {"delegated_authority_partner": "delegated_authority_partner"}

    response = client.patch(f"/api/silver/runs/{run['id']}/spread", json={
        "table_name": "ext_pc0002_data_agg", "columns": table["spread"]["columns"], "measure": "premium",
        "dimension": "product_line_name"})
    assert response.status_code == 200, response.text
    quality = response.json()["tables"][0]["quality"]
    assert quality["rows_loaded"] == 30 and quality["rollup_rows"] == 2

    _approve(client, response.json())
    rows = _q(config, "SELECT delegated_authority_partner, product_line_name, premium, reporting_period, record_grain, "
                      "agg_data_source FROM {s}.silver_aggregate WHERE source_system = 'pc0002'")
    assert len(rows) == 30 and "TOTALS" not in {row[0] for row in rows}
    assert {row[1] for row in rows} == {"Casualty Treaty", "Property Treaty", "Workers Comp"}
    assert {row[3] for row in rows} == {"2026-01", "2026-03"}
    assert {(row[4], row[5]) for row in rows} == {("AS_REPORTED", "bronze_aggregate")}
    total = round(float(sum(row[2] for row in rows)), 2)
    assert total == round(2 * (9928.42 + 5593.17 + 443.81 + 2643.97 + 1375.23 + 207.54 + 4036.57 + 6538.9 + 925.17
                               + 3640.43 + 242.84), 2)
    # Rebuilding the roll-up from the transaction table keeps the profit center's own rows.
    from api import db
    from api.services import silver_service

    with db.connection() as conn:
        silver_service._rebuild_aggregate(conn, ["pc0002"])
    assert _q(config, "SELECT count(*) FROM {s}.silver_aggregate WHERE source_system = 'pc0002'")[0][0] == 30
