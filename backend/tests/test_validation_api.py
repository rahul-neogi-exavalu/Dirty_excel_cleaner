"""The Validate step through the HTTP API, without a database: what the database would say
(the bronze column mapping, the control table) is stubbed."""

import io
import sys
import threading
import time
from datetime import date
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402
from sse_client import read_events, validate  # noqa: E402

HEADERS = ["AccountingEffectiveDate", "CommissionPct", "GrossCommissionAmount", "InsuranceCompany Name",
           "MarketProvider", "Pol_Ef_Dt", "PolicyNumber", "Premium", "Producer/Agency Name",
           "ProducerCommissionAmount", "ProducerCommissionPct", "Revenue", "Transaction Detail",
           "TransactionEffectiveDate"]


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from api import config

    work = tmp_path_factory.mktemp("workspace")
    config.WORK_DIR = work
    config.UPLOAD_DIR = work / "uploads"
    config.JOB_DIR = work / "jobs"
    from api.main import app

    return TestClient(app)


@pytest.fixture
def context(monkeypatch):
    """The database's answers: one saved mapping for PC0101 (Pol_Ef_Dt is its PED), and
    whatever each test puts in ``loaded``."""
    from api import config
    from api.services import validation_service

    monkeypatch.setattr(config, "DB_CONFIGURED", True)
    monkeypatch.setattr(config, "INGEST_ENABLED", True)
    monkeypatch.setattr(config, "AI_ENABLED", False)
    monkeypatch.setattr(config, "WORD2VEC_PATH", "")
    found = validation_service.Context(
        divisions=[("AH Programs", "No", "0101")],
        mapping_rows=[("PC0101", "Pol_Ef_Dt", None, "policy_effective_date")],
    )
    monkeypatch.setattr(validation_service, "_read", lambda names, conn=None, tracker=None: found)
    return found


DATES = ("AccountingEffectiveDate", "Pol_Ef_Dt", "TransactionEffectiveDate")


def _fill(sheet, months, headers=HEADERS, blank=(), ped=None):
    """``blank``: the columns left empty on the first row."""
    sheet.append(headers)
    n = 0
    for month in months:
        for day in (3, 17):
            n += 1
            row = {
                "AccountingEffectiveDate": date(2026, month, day),
                "CommissionPct": 0.1, "GrossCommissionAmount": 100.5 + n, "InsuranceCompany Name": "Harrier",
                "MarketProvider": "Market A", "Pol_Ef_Dt": ped or date(2026, month, 1), "PolicyNumber": f"POL-{n:04d}",
                "Premium": 1000 + n, "Producer/Agency Name": "Brown & Brown", "ProducerCommissionAmount": 50 + n,
                "ProducerCommissionPct": 0.05, "Revenue": 40 + n, "Transaction Detail": "New",
                "TransactionEffectiveDate": date(2026, month, day), "Notes": f"note {n}",
            }
            if n == 1:
                row.update(dict.fromkeys(blank))
            sheet.append([row.get(name) for name in headers])


def _save(book):
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _book(months, headers=HEADERS, blank_aed=False, ped=None):
    book = openpyxl.Workbook()
    book.active.title = "Data"
    _fill(book.active, months, headers, ("AccountingEffectiveDate",) if blank_aed else (), ped)
    return _save(book)


def _sheets(*sheets):
    """A workbook of several sheets: (title, months, headers, blank) each. Different headers
    keep them apart (the cleaner stacks sheets whose headers are the same)."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, months, headers, blank in sheets:
        _fill(book.create_sheet(title), months, headers, blank)
    return _save(book)


def _clean(client, name, content):
    upload = client.post("/api/workbooks", files={"file": (name, content)}).json()
    job = client.post("/api/jobs", json={"workbook_id": upload["id"], "sheets": [s["name"] for s in upload["sheets"]]}).json()
    for _ in range(300):
        status = client.get(f"/api/jobs/{job['id']}").json()
        if status["status"] not in ("queued", "running"):
            break
        time.sleep(0.05)
    assert status["status"] == "succeeded", status
    return job["id"]


def _validate(client, name, content):
    data = validate(client, [_clean(client, name, content)])
    return data, data["outputs"][0]


def _required(output, name):
    return next(item for item in output["required"] if item["name"] == name)


def _by_sheet(data):
    return {output["sheet_name"]: output for output in data["outputs"]}


def _check(output, check_id):
    return next(check for check in output["checks"] if check["id"] == check_id)


def _patch(client, data, output, change):
    response = client.patch(f"/api/validations/{data['id']}/outputs/{output['key']}", json=change)
    assert response.status_code == 200, response.text
    return response.json()


def test_a_year_to_date_file_with_every_column_is_ready(client, context):
    data, output = _validate(client, "PC0101_Report_07132026.xlsx", _book(range(1, 7)))
    file = data["files"][0]
    assert (file["pc_id"], file["source_system"], file["file_received_date"]) == ("PC0101", "EXT_PC0101", "2026-07-13")
    assert file["division_name"] == "AH Programs"
    assert output["missing"] == [] and all(item["column"] for item in output["required"])
    # The bronze column mapping (seeded from the DRT mapping) knows Pol_Ef_Dt for PC0101.
    assert _required(output, "policy_effective_date")["vote"]["method"] == "saved"
    assert _required(output, "premium")["vote"]["method"] == "exact"
    assert (output["date_detail"], output["period_type"]) == ("AED", "YTD")
    assert (output["reporting_start_date"], output["reporting_end_date"]) == ("2026-01-01", "2026-06-30")
    assert output["action"] == "INSERT" and output["verdict"] == "ready" and output["fitness"] == 100
    assert data["counts"]["ready"] == 1


def test_one_empty_aed_hands_the_dates_to_ted_and_a_non_ytd_span_is_flagged(client, context):
    data, output = _validate(client, "PC0101_Partial_07132026.xlsx", _book(range(2, 7), blank_aed=True))
    assert output["dates"]["AED"]["complete"] is False and output["dates"]["TED"]["complete"] is True
    assert output["date_detail"] == "TED" and output["period_type"] is None
    assert output["verdict"] == "flagged" and output["action"] == "REJECTED" and output["flag"]
    # AED misses a row: it cannot be picked to decide the dates.
    response = client.patch(f"/api/validations/{data['id']}/outputs/{output['key']}", json={"date_role": "AED"})
    assert response.status_code == 422 and response.json()["error"]["code"] == "date_role_incomplete"


def _dated_book(rows):
    """One sheet whose AED / TED / PED are given per row (every other column filled)."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(HEADERS)
    for n, (aed, ted, ped) in enumerate(rows, 1):
        values = {"AccountingEffectiveDate": aed, "TransactionEffectiveDate": ted, "Pol_Ef_Dt": ped,
                  "CommissionPct": 0.1, "GrossCommissionAmount": 100.5 + n, "InsuranceCompany Name": "Harrier",
                  "MarketProvider": "Market A", "PolicyNumber": f"POL-{n:04d}", "Premium": 1000 + n,
                  "Producer/Agency Name": "Brown & Brown", "ProducerCommissionAmount": 50 + n,
                  "ProducerCommissionPct": 0.05, "Revenue": 40 + n, "Transaction Detail": "New"}
        sheet.append([values.get(name) for name in HEADERS])
    return _save(book)


def test_the_reviewer_picks_the_date_column_and_it_is_the_date_detail(client, context):
    # AED runs Feb-Jun (neither YTD nor monthly), TED Jan-Jun (YTD), PED is one month of 2025.
    rows = [(date(2026, max(month, 2), 10), date(2026, month, 5), date(2025, 7, 1)) for month in range(1, 7)]
    data, output = _validate(client, "PC0101_Picked_07132026.xlsx", _dated_book(rows))
    assert output["date_detail"] == "AED" and output["verdict"] == "flagged"
    roles = output["date_roles"]
    assert roles["AED"]["period_type"] is None and roles["AED"]["flag"]
    assert (roles["TED"]["start"], roles["TED"]["end"], roles["TED"]["period_type"]) == ("2026-01-01", "2026-06-30", "YTD")
    assert roles["PED"]["period_type"] == "MONTHLY"

    # Picking TED: its range decides and TED is recorded, not a blank.
    picked = _patch(client, data, output, {"date_role": "TED"})["outputs"][0]
    assert (picked["date_detail"], picked["date_role"], picked["origin"]) == ("TED", "TED", "picked")
    assert (picked["reporting_start_date"], picked["reporting_end_date"], picked["period_type"]) == (
        "2026-01-01", "2026-06-30", "YTD")
    assert picked["verdict"] == "ready" and picked["month_column"] == _required(picked, "transaction_effective_date")["column"]
    assert "picked by the reviewer" in _check(picked, "dates")["detail"]

    # Typed dates that are TED's range: TED. Dates no column spans: blank.
    typed = _patch(client, data, output, {"reporting_start_date": "2026-01", "reporting_end_date": "2026-06"})["outputs"][0]
    assert (typed["date_detail"], typed["origin"], typed["date_role"]) == ("TED", "entered", None)
    made_up = _patch(client, data, output, {"reporting_start_date": "2026-01", "reporting_end_date": "2026-05"})["outputs"][0]
    assert made_up["date_detail"] is None and made_up["origin"] == "entered"
    assert "no date column's range" in _check(made_up, "dates")["detail"]

    # Back to the priority's column.
    _patch(client, data, output, {"reporting_start_date": None, "reporting_end_date": None})
    back = _patch(client, data, output, {"date_role": None})["outputs"][0]
    assert back["date_detail"] == "AED" and back["origin"] == "AED"


def test_the_reviewer_corrects_flagged_dates(client, context):
    data, output = _validate(client, "PC0101_FebJun_07132026.xlsx", _book(range(2, 7)))
    assert output["verdict"] == "flagged"
    fixed = client.patch(f"/api/validations/{data['id']}/outputs/{output['key']}",
                         json={"reporting_start_date": "2026-01", "reporting_end_date": "2026-06"}).json()["outputs"][0]
    assert fixed["verdict"] == "ready" and fixed["period_type"] == "YTD" and fixed["date_detail"] is None
    assert fixed["origin"] == "entered"


def test_a_missing_required_column_rejects_the_file(client, context):
    headers = [name for name in HEADERS if name != "Revenue"]
    data, output = _validate(client, "PC0101_NoRevenue_07132026.xlsx", _book(range(1, 7), headers=headers))
    assert output["missing"] == ["Revenue"] and output["verdict"] == "rejected" and output["action"] == "REJECTED"
    # Mapping another column to it would be the reviewer's call; clearing a found one rejects too.
    key = output["key"]
    cleared = client.patch(f"/api/validations/{data['id']}/outputs/{key}",
                           json={"mapping": {"premium": None}}).json()["outputs"][0]
    assert set(cleared["missing"]) == {"Revenue", "Premium"}
    assert _required(cleared, "premium")["reviewer"] is True


# The year-to-date file PC0101 sent first, as the control table has it: its bronze columns.
H1_COLUMNS = ("accountingeffectivedate", "commissionpct", "grosscommissionamount", "insurancecompany_name",
              "marketprovider", "pol_ef_dt", "policynumber", "premium", "producer_agency_name",
              "producercommissionamount", "producercommissionpct", "revenue", "transaction_detail",
              "transactioneffectivedate")


def _h1(columns=H1_COLUMNS):
    from ahi_bronze.validation import Loaded

    return Loaded(5, "PC0101_H1_07132026.xlsx", "YTD", date(2026, 1, 1), date(2026, 6, 30), rows=120, columns=14,
                  months={"2026-06": {"rows": 20, "AED": 100.0, "PED": 100.0, "TED": 95.0}}, column_names=columns)


def test_a_month_already_in_bronze_is_flagged_and_rejected(client, context):
    context.loaded = {"EXT_PC0101": [_h1()]}
    # June again, in a file of another name: the profit center in the name ties it to the H1 file.
    data, output = _validate(client, "PC0101_June_07132026.xlsx", _book([6]))
    assert output["period_type"] == "MONTHLY" and output["action"] == "REJECTED"
    assert output["verdict"] == "flagged" and output["options"] == []
    assert output["reasons"][0].startswith("Jun 2026 is already loaded from PC0101_H1_07132026.xlsx")
    compare = output["compare"]
    assert compare["month"] == "2026-06" and compare["earlier"][0]["rows"] == 20 and compare["this"]["rows"] == 2


def test_a_month_after_the_year_to_date_file_is_appended_and_its_schema_compared(client, context):
    context.loaded = {"EXT_PC0101": [_h1()]}
    # July, in a file of its own name, with one column the H1 file did not have (Notes).
    data, july = _validate(client, "PC0101_July_08142026.xlsx", _book([7], headers=HEADERS + ["Notes"]))
    assert july["action"] == "APPEND" and july["verdict"] == "ready"
    assert [item["file_name"] for item in july["appends_to"]] == ["PC0101_H1_07132026.xlsx"]
    schema = july["schema"]
    assert schema["against"] == "PC0101_H1_07132026.xlsx" and schema["kind"] == "evolved"
    assert schema["added"] == ["notes"] and schema["missing"] == []
    # The same columns as H1: identical.
    _, same = _validate(client, "PC0101_July_08142026.xlsx", _book([7]))
    assert same["schema"]["kind"] == "identical"


# --- the profit center's own aggregates --------------------------------------------------


def _summary_book():
    """A premium summary a month to a sheet, spaced for print like PC2030's: a title, the
    header two blank rows above the partners, a blank row after each, spacer columns, a
    TOTALS column and a TOTALS row -- and a sheet saying no premium was written."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    partners = [("Accident Fund", 9928.42, 5593.17, 443.81), ("CompSource Mutual", 2643.97, 1375.23, 207.54),
                ("Convex UK", 0, 4036.57, 0), ("Hastings", 6538.9, 0, 925.17), ("Tokio Marine", 0, 3640.43, 242.84)]
    for month in ("January 2026", "February 2026", "March 2026"):
        sheet = book.create_sheet(month)
        if month.startswith("February"):
            sheet.append([None, "NO PREMIUM WRITTEN IN FEBRUARY"])
            continue
        for row in (["Premium Summary by Partner"], [month], [], [],
                    ["Delegated Authority Partner", "Casualty Treaty", None, "Property Treaty", None, "Workers Comp",
                     None, "TOTALS"], [], []):
            sheet.append(row)
        for name, casualty, prop, workers in partners:
            sheet.append([name, casualty, None, prop, None, workers, None, round(casualty + prop + workers, 2)])
            sheet.append([])
        sums = [round(sum(item[index] for item in partners), 2) for index in (1, 2, 3)]
        sheet.append(["TOTALS", sums[0], None, sums[1], None, sums[2], None, round(sum(sums), 2)])
    return _save(book)


def test_a_summary_the_profit_center_aggregated_is_detected_and_dated_by_its_sheets(client, context):
    data, output = _validate(client, "PC0101_Premium Summary January 1 2026 to March 31 2026_04152026.xlsx",
                             _summary_book())
    assert output["aggregated"] is True
    found = output["grain"]
    assert found["total_column"] == "totals" and found["total_rows"] == 2 and found["dimensions"] == [
        "delegated_authority_partner"]
    # Not held to the transaction columns; dated by the month each sheet names.
    assert output["missing"] == [] and output["origin"] == "sheets" and output["date_detail"] is None
    assert output["sheet_months"] == {"January 2026": "2026-01", "March 2026": "2026-03"}
    assert (output["reporting_start_date"], output["reporting_end_date"], output["period_type"]) == (
        "2026-01-01", "2026-03-31", "YTD")
    assert output["action"] == "INSERT" and output["verdict"] == "ready"
    assert set(output["months"]) == {"2026-01", "2026-03"}

    # The reviewer may say otherwise: as transactions it misses the business's columns.
    as_rows = _patch(client, data, output, {"grain": "transaction"})["outputs"][0]
    assert as_rows["aggregated"] is False and as_rows["verdict"] == "rejected" and "PolicyNumber" in as_rows["missing"]
    back = _patch(client, data, output, {"grain": None})["outputs"][0]
    assert back["aggregated"] is True and back["grain"]["override"] is None


def test_aggregates_are_checked_against_the_profit_centers_aggregates_only(client, context):
    from ahi_bronze.validation import Loaded

    # A transaction file of Jan-Mar is loaded; the summary of the same months is no revision of it.
    context.loaded = {"EXT_PC0101": [Loaded(9, "PC0101_Txn_04152026.xlsx", "YTD", date(2026, 1, 1),
                                            date(2026, 3, 31), rows=40)]}
    _, output = _validate(client, "PC0101_Premium Summary_04152026.xlsx", _summary_book())
    assert output["aggregated"] and output["action"] == "INSERT" and output["revisable"] == []
    context.loaded = {"EXT_PC0101": [Loaded(9, "PC0101_Premium Summary_03152026.xlsx", "YTD", date(2026, 1, 1),
                                            date(2026, 3, 31), rows=12, aggregated=True)]}
    _, again = _validate(client, "PC0101_Premium Summary_04152026.xlsx", _summary_book())
    assert again["action"] == "DECIDE" and again["options"] == ["revise", "companion", "reject"]


def test_a_revision_names_the_loaded_file_it_replaces(client, context):
    from ahi_bronze.validation import Loaded

    context.loaded = {"EXT_PC0101": [
        Loaded(5, "PC0101_Report_06302026.xlsx", "YTD", date(2026, 1, 1), date(2026, 6, 30), rows=12, columns=14),
        Loaded(6, "PC0101_July_08012026.xlsx", "MONTHLY", date(2026, 7, 1), date(2026, 7, 31), rows=2, columns=14),
    ]}
    data, output = _validate(client, "PC0101_Report_07132026.xlsx", _book(range(1, 7)))
    # The same reporting dates as a loaded file: revision, companion or reject.
    assert output["action"] == "DECIDE" and output["options"] == ["revise", "companion", "reject"]
    assert output["verdict"] == "needs_input"
    assert [item["file_name"] for item in output["revisable"]] == ["PC0101_Report_06302026.xlsx"]

    url = f"/api/validations/{data['id']}/outputs/{output['key']}"
    # The file must be named, and named exactly as the control table has it.
    unnamed = client.patch(url, json={"choice": "revise"})
    assert unnamed.status_code == 422 and unnamed.json()["error"]["code"] == "revised_file_required"
    assert unnamed.json()["error"]["advice"] == "Name one of: PC0101_Report_06302026.xlsx."
    unknown = client.patch(url, json={"choice": "revise", "replaces_file": "PC0101_Report.xlsx"})
    assert unknown.status_code == 422 and unknown.json()["error"]["code"] == "unknown_revised_file"
    assert "is not a loaded file of PC0101 for Jan 2026 – Jun 2026" in unknown.json()["error"]["message"]
    # July shares no month with Jan-Jun: not a file this one can revise.
    july = client.patch(url, json={"choice": "revise", "replaces_file": "PC0101_July_08012026.xlsx"})
    assert july.status_code == 422 and july.json()["error"]["code"] == "unknown_revised_file"

    revised = _patch(client, data, output, {"choice": "revise", "replaces_file": "PC0101_Report_06302026.xlsx"})["outputs"][0]
    assert revised["action"] == "INSERT" and revised["verdict"] == "ready" and revised["confirm"]
    assert revised["file_replaced"] == revised["replaces_file"] == "PC0101_Report_06302026.xlsx"
    assert any("Keeps the later PC0101_July_08012026.xlsx" in reason for reason in revised["reasons"])
    # Undoing the choice asks again.
    undone = _patch(client, data, output, {"choice": None})["outputs"][0]
    assert undone["action"] == "DECIDE" and undone["replaces_file"] is None


def test_no_profit_center_or_received_date_needs_input(client, context):
    data, output = _validate(client, "Report_without_codes.xlsx", _book(range(1, 7)))
    assert output["verdict"] == "needs_input" and data["files"][0]["pc_id"] is None
    job = data["files"][0]["job_id"]
    after = client.patch(f"/api/validations/{data['id']}/files/{job}", json={"pc_id": "101"}).json()
    assert after["files"][0]["pc_id"] == "PC0101"
    output = after["outputs"][0]
    assert output["verdict"] == "needs_input" and "Enter the file received date." in output["needs"]
    done = client.patch(f"/api/validations/{data['id']}/files/{job}", json={"file_received_date": "2026-07-13"}).json()
    assert done["outputs"][0]["verdict"] == "ready"
    bad = client.patch(f"/api/validations/{data['id']}/files/{job}", json={"pc_id": "abc"})
    assert bad.status_code == 422


def test_either_column_of_a_one_of_pair_is_enough(client, context):
    # CommissionPct without GrossCommissionAmount; ProducerCommissionAmount without ProducerCommissionPct.
    headers = [name for name in HEADERS if name not in ("GrossCommissionAmount", "ProducerCommissionPct")]
    _, output = _validate(client, "PC0101_OneOfEach_07132026.xlsx", _book(range(1, 7), headers=headers))
    assert output["missing"] == [] and output["verdict"] == "ready" and output["fitness"] == 100
    assert _required(output, "gross_commission_amount")["column"] is None
    assert _required(output, "commission_pct")["one_of"] == ["gross_commission_amount"]
    assert _required(output, "producer_commission_pct")["one_of"] == ["producer_commission_amount"]
    assert _required(output, "revenue")["one_of"] == []
    assert _check(output, "columns")["detail"] == "12 of 12 found"


def test_neither_column_of_a_one_of_pair_rejects_the_file(client, context):
    headers = [name for name in HEADERS if name not in ("CommissionPct", "GrossCommissionAmount")]
    _, output = _validate(client, "PC0101_NoCommission_07132026.xlsx", _book(range(1, 7), headers=headers))
    assert output["missing"] == ["CommissionPct or GrossCommissionAmount"]
    assert output["verdict"] == "rejected" and output["action"] == "REJECTED"
    assert "Missing required column: CommissionPct or GrossCommissionAmount." in output["reasons"]


def test_clearing_one_column_of_a_found_pair_leaves_the_file_fit(client, context):
    data, output = _validate(client, "PC0101_Both_07132026.xlsx", _book(range(1, 7)))
    assert output["verdict"] == "ready"
    cleared = _patch(client, data, output, {"mapping": {"commission_pct": None}})["outputs"][0]
    assert cleared["missing"] == [] and cleared["verdict"] == "ready"
    both = _patch(client, data, output, {"mapping": {"gross_commission_amount": None}})["outputs"][0]
    assert both["missing"] == ["CommissionPct or GrossCommissionAmount"] and both["verdict"] == "rejected"


def test_a_sheet_without_the_required_columns_is_left_out_and_the_valid_sheet_goes_on(client, context):
    no_revenue = [name for name in HEADERS if name != "Revenue"]
    data, _ = _validate(client, "PC0101_TwoSheets_07132026.xlsx", _sheets(
        ("Fit", range(1, 7), HEADERS, ()), ("NoRevenue", range(1, 7), no_revenue, ())))
    sheets = _by_sheet(data)
    assert set(sheets) == {"Fit", "NoRevenue"}
    left_out = sheets["NoRevenue"]
    assert left_out["missing"] == ["Revenue"] and left_out["verdict"] == "rejected"
    assert any("Left out on its own" in warning for warning in left_out["warnings"])
    # The valid sheet is not dragged down with it.
    fit = sheets["Fit"]
    assert fit["missing"] == [] and fit["verdict"] == "ready" and fit["action"] == "INSERT"
    assert _check(fit, "file") == {"id": "file", "label": "Every sheet of the file", "ok": True,
                                   "detail": "NoRevenue left out (not data)"}
    assert fit["fitness"] == 100 and data["counts"]["ready"] == 1 and data["counts"]["rejected"] == 1
    # Mapping a column to Revenue makes that sheet fit too.
    premium = next(column["original"] for column in left_out["columns"] if column["header"] == "Premium")
    after = _by_sheet(_patch(client, data, left_out, {"mapping": {"revenue": premium}}))
    assert {name: output["verdict"] for name, output in after.items()} == {"Fit": "ready", "NoRevenue": "ready"}
    assert _check(after["Fit"], "file")["detail"] == "The other 1 sheet fit"


def test_a_data_sheet_unfit_for_another_reason_still_rejects_the_file(client, context):
    # Both sheets carry every required column; FebJun's dates are neither YTD nor monthly.
    data, _ = _validate(client, "PC0101_Flagged_07132026.xlsx", _sheets(
        ("Fit", range(1, 7), HEADERS, ()), ("FebJun", range(2, 7), HEADERS + ["Notes"], ())))
    sheets = _by_sheet(data)
    assert sheets["FebJun"]["verdict"] == "flagged" and sheets["FebJun"]["missing"] == []
    fit = sheets["Fit"]
    assert fit["verdict"] == "rejected" and fit["action"] == "REJECTED"
    assert "every data sheet of a file must be fit" in fit["reasons"][0] and "FebJun" in fit["reasons"][0]
    assert _check(fit, "file") == {"id": "file", "label": "Every sheet of the file", "ok": False,
                                   "detail": "Rejected with FebJun"}


def test_a_sheet_waiting_for_the_reviewer_holds_the_file_back(client, context):
    data, _ = _validate(client, "PC0101_Undated_07132026.xlsx", _sheets(
        ("Fit", range(1, 7), HEADERS, ()), ("Undated", range(1, 7), HEADERS + ["Notes"], DATES)))
    sheets = _by_sheet(data)
    assert sheets["Undated"]["verdict"] == "needs_input" and sheets["Undated"]["reporting_start_date"] is None
    fit = sheets["Fit"]
    assert fit["verdict"] == "needs_input" and fit["action"] == "INSERT"
    assert fit["needs"] == ["Finish Undated first: a file's sheets are staged together."]
    assert _check(fit, "file")["detail"] == "Waiting on Undated"
    after = _by_sheet(_patch(client, data, sheets["Undated"],
                             {"reporting_start_date": "2026-01", "reporting_end_date": "2026-06"}))
    assert after["Fit"]["verdict"] == "ready" and after["Undated"]["verdict"] == "ready"
    # The reviewer rejecting one sheet rejects the other.
    rejected = _by_sheet(_patch(client, data, after["Fit"], {"choice": "reject"}))
    assert rejected["Fit"]["reasons"] == ["Rejected by the reviewer."]
    assert rejected["Undated"]["verdict"] == "rejected" and "Fit (Rejected by the reviewer)" in rejected["Undated"]["reasons"][0]
    assert any("so are the file's other 1 sheet" in warning for warning in rejected["Fit"]["warnings"])


def test_staging_one_sheet_stages_the_whole_file(client, context, monkeypatch):
    from contextlib import contextmanager

    from api.services import validation_service

    class Conn:
        def execute(self, *args, **kwargs):
            return self

    @contextmanager
    def connection():
        yield Conn()

    staged = []

    def stage_one(conn, session, output, user_name, companion_of=None):
        staged.append(output.key)
        return {"control_id": len(staged), "staging_table": f"stg_{len(staged)}",
                "processing_action": output.decision.action, "seeded": False}

    monkeypatch.setattr(validation_service.db, "connection", connection)
    monkeypatch.setattr(validation_service, "_stage_one", stage_one)
    data, _ = _validate(client, "PC0101_Both_07132026.xlsx", _sheets(
        ("Jan", range(1, 7), HEADERS, ()), ("Feb", range(1, 7), HEADERS + ["Notes"], ())))
    keys = [output["key"] for output in data["outputs"]]
    assert len(keys) == 2
    response = client.post(f"/api/validations/{data['id']}/stage", json={"keys": keys[:1]})
    assert response.status_code == 200, response.text
    assert sorted(staged) == sorted(keys) and response.json()["staged"] == 2


# --- companion files: a list that came with the data ------------------------------------


BROKERS = [("Brown & Brown", "Market A", "East"), ("Acme", "Market B", "West"), ("Lockton", "Market A", "North"),
           ("Marsh", "Market C", "South"), ("Aon", "Market B", "East"), ("Gallagher", "Market C", "West"),
           ("Hub", "Market A", "North"), ("Alliant", "Market B", "South")]


def _broker_list(rows=BROKERS):
    """A broker list: who each broker is, and the market provider the data file lacks."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Brokers"
    sheet.append(["Account_Name", "MarketProvider", "Region", "Agents"])
    for index, row in enumerate(rows, 1):
        sheet.append(list(row) + [index * 3])
    return _save(book)


def _pair(client):
    """PC0101's transactions (no MarketProvider, a Producer_name instead) and its broker
    list, received together on 8 Aug 2026, validated in one batch."""
    headers = [name for name in HEADERS if name != "MarketProvider"]
    data = validate(client, [_clean(client, "PC0101_Transactional_08082026.xlsx", _book(range(1, 7), headers=headers)),
                             _clean(client, "PC0101_BrokerList_08082026.xlsx", _broker_list())])
    by_file = {file["job_id"]: file["file_name"] for file in data["files"]}
    outputs = {by_file[output["job_id"]].split("_")[1]: output for output in data["outputs"]}
    return data, outputs["Transactional"], outputs["BrokerList"]


def test_a_list_that_came_with_the_data_is_its_companion(client, context):
    data, transactions, brokers = _pair(client)
    # On their own each lacks a required column; the broker list may be the other's companion.
    assert transactions["missing"] == ["MarketProvider"] and transactions["verdict"] == "rejected"
    assert brokers["verdict"] == "rejected" and brokers["options"] == ["companion"]
    assert [partner["ref"] for partner in brokers["partners"]] == [transactions["key"]]
    assert brokers["partners"][0]["file_name"] == "PC0101_Transactional_08082026.xlsx"

    url = f"/api/validations/{data['id']}/outputs/{brokers['key']}"
    wrong = client.patch(url, json={"choice": "companion", "companion_of": "nothing"})
    assert wrong.status_code == 422 and wrong.json()["error"]["code"] == "unknown_companion_of"
    assert "PC0101_Transactional_08082026.xlsx" in wrong.json()["error"]["advice"]

    after = _patch(client, data, brokers, {"choice": "companion", "companion_of": transactions["key"]})
    outputs = {output["key"]: output for output in after["outputs"]}
    brokers, transactions = outputs[brokers["key"]], outputs[transactions["key"]]
    # Checked together: the broker list carries MarketProvider, so the pair is fit.
    assert transactions["missing"] == [] and transactions["verdict"] == "ready" and transactions["action"] == "INSERT"
    assert transactions["companions"] == ["PC0101_BrokerList_08082026.xlsx"]
    assert transactions["pair"] == "PC0101_Transactional_08082026.xlsx + PC0101_BrokerList_08082026.xlsx"
    # The companion reports what the data file reports.
    assert brokers["action"] == "APPEND" and brokers["verdict"] == "ready" and brokers["origin"] == "companion"
    assert (brokers["reporting_start_date"], brokers["reporting_end_date"], brokers["period_type"], brokers["date_detail"]) == (
        "2026-01-01", "2026-06-30", "YTD", "AED")
    assert brokers["companion"]["ref"] == transactions["key"] and "Companion of" in brokers["reasons"][0]
    # The data file cannot become a companion while one points at it.
    taken = client.patch(f"/api/validations/{data['id']}/outputs/{transactions['key']}",
                         json={"choice": "companion", "companion_of": brokers["key"]})
    assert taken.status_code == 422


def test_a_companion_is_staged_with_the_file_it_came_with(client, context, monkeypatch):
    from contextlib import contextmanager

    from api.services import validation_service

    class Conn:
        def execute(self, *args, **kwargs):
            return self

    @contextmanager
    def connection():
        yield Conn()

    staged = []

    def stage_one(conn, session, output, user_name, companion_of=None):
        staged.append((output.key, companion_of))
        return {"control_id": 40 + len(staged), "staging_table": f"t_{len(staged)}",
                "processing_action": output.decision.action, "seeded": False}

    data, transactions, brokers = _pair(client)
    _patch(client, data, brokers, {"choice": "companion", "companion_of": transactions["key"]})
    monkeypatch.setattr(validation_service.db, "connection", connection)
    monkeypatch.setattr(validation_service, "_stage_one", stage_one)
    # Staging the broker list stages the data file too, first, and points at its control row.
    response = client.post(f"/api/validations/{data['id']}/stage", json={"keys": [brokers["key"]]})
    assert response.status_code == 200, response.text
    assert staged == [(transactions["key"], None), (brokers["key"], 41)]


def test_a_companion_of_a_loaded_file_takes_its_dates(client, context):
    from ahi_bronze.validation import Loaded
    from api.services.validation_service import required

    # The loaded data file carries every required column but MarketProvider.
    carried = frozenset(item.name for item in required() if item.name != "market_provider")
    context.loaded = {"EXT_PC0101": [Loaded(7, "PC0101_Transactional_08082026.xlsx", "YTD", date(2026, 1, 1),
                                            date(2026, 6, 30), rows=12, received=date(2026, 8, 8), date_detail="TED",
                                            present=carried)]}
    data, brokers = _validate(client, "PC0101_BrokerList_08082026.xlsx", _broker_list())
    assert [partner["ref"] for partner in brokers["partners"]] == ["control:7"]
    linked = _patch(client, data, brokers, {"choice": "companion", "companion_of": "7"})["outputs"][0]
    assert linked["companion_of"] == "control:7" and linked["action"] == "APPEND"
    assert (linked["reporting_start_date"], linked["date_detail"], linked["period_type"]) == ("2026-01-01", "TED", "YTD")


# --- the progress of a validation, streamed -------------------------------------------


def _start(client, name, content):
    response = client.post("/api/validations", json={"job_ids": [_clean(client, name, content)]})
    assert response.status_code == 202, response.text
    return response.json()


def _ai(monkeypatch, answer):
    """The AI matcher replaced by ``answer`` (columns, catalog) -> {}: AI matching runs."""
    from api.services import column_matching

    monkeypatch.setattr(column_matching, "matchers",
                        lambda precedents=(), vectors_progress=None: (None, answer, []))


def test_the_stream_walks_each_step_to_the_review(client, context, monkeypatch):
    def slow_ai(columns, catalog):
        time.sleep(0.4)  # long enough for the stream to show the AI at work
        return {}

    _ai(monkeypatch, slow_ai)
    started = _start(client, "PC0101_Progress_07132026.xlsx", _sheets(
        ("Jan", range(1, 7), HEADERS, ()), ("Feb", range(1, 7), HEADERS + ["Notes"], ())))
    assert [phase["id"] for phase in started["progress"]["phases"]] == ["connect", "read", "matchers", "match", "check"]
    assert started["progress"]["counters"]["tables_matched"] == {"done": 0, "total": 2}
    events = read_events(client, f"/api/validations/{started['id']}/events")
    names = [name for name, _ in events]
    assert names[-1] == "done" and set(names[:-1]) == {"progress"}
    states = [data for name, data in events if name == "progress"] + [events[-1][1]["progress"]]
    fractions = [state["fraction"] for state in states]
    assert fractions == sorted(fractions) and states[-1]["percent"] == 100
    order = ["connect", "read", "matchers", "match", "check"]
    steps = [order.index(state["phase"]) for state in states[:-1] if state["phase"]]
    assert steps == sorted(steps) and order.index("match") in steps
    asked = [state for state in states if "asking the AI to read" in (state["detail"] or "")]
    assert asked and asked[0]["detail"].startswith("PC0101_Progress_07132026.xlsx · ")
    assert 0 < asked[0]["fraction"] < 1 and asked[0]["counters"]["tables_matched"]["done"] < 2
    final = states[-1]
    assert final["counters"]["tables_matched"] == {"done": 2, "total": 2}
    assert final["counters"]["rows_checked"] == {"done": 24, "total": 24}
    assert all(phase["state"] == "done" for phase in final["phases"])
    review = events[-1][1]["result"]
    assert review["status"] == "ready" and len(review["outputs"]) == 2
    # The review is also what GET answers now.
    assert client.get(f"/api/validations/{started['id']}").json()["outputs"] == review["outputs"]


def test_a_running_validation_shows_its_progress_and_waits_for_edits(client, context, monkeypatch):
    release = threading.Event()

    def held_ai(columns, catalog):
        assert release.wait(10)
        return {}

    _ai(monkeypatch, held_ai)
    started = _start(client, "PC0101_Held_07132026.xlsx", _book(range(1, 7)))
    url = f"/api/validations/{started['id']}"
    try:
        for _ in range(200):
            running = client.get(url).json()
            if "asking the AI" in (running["progress"]["detail"] or ""):
                break
            time.sleep(0.05)
        assert running["status"] == "running" and running["progress"]["phase"] == "match"
        assert running["progress"]["detail"] == "PC0101_Held_07132026.xlsx: asking the AI to read 14 columns"
        assert "outputs" not in running and 0 < running["progress"]["percent"] < 100
        busy = client.post(f"{url}/stage", json={})
        assert busy.status_code == 409 and busy.json()["error"]["message"] == "This validation is still running."
    finally:
        release.set()
    events = read_events(client, f"{url}/events")
    assert events[-1][0] == "done" and client.get(url).json()["status"] == "ready"


def test_a_failed_validation_says_why_on_the_stream(client, context, monkeypatch):
    from api.errors import ApiError
    from api.services import silver_service

    def broken():
        raise ApiError(500, "silver_catalog", "The Silver column list could not be read.", "Check silver_columns.csv.")

    monkeypatch.setattr(silver_service, "catalog", broken)
    started = _start(client, "PC0101_Broken_07132026.xlsx", _book(range(1, 7)))
    events = read_events(client, f"/api/validations/{started['id']}/events")
    name, data = events[-1]
    assert name == "failed" and data["error"]["code"] == "silver_catalog"
    assert data["error"]["advice"] == "Check silver_columns.csv." and data["progress"]["status"] == "failed"
    failed = client.get(f"/api/validations/{started['id']}").json()
    assert failed["status"] == "failed" and failed["error"]["message"] == "The Silver column list could not be read."
    patched = client.patch(f"/api/validations/{started['id']}/files/x", json={"pc_id": "PC0101"})
    assert patched.status_code == 409


def test_an_unexpected_error_is_reported_not_swallowed(client, context, monkeypatch):
    from api.services import validation_service

    def crash(*args, **kwargs):
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(validation_service, "_evaluate", crash)
    started = _start(client, "PC0101_Crash_07132026.xlsx", _book(range(1, 7)))
    name, data = read_events(client, f"/api/validations/{started['id']}/events")[-1]
    assert name == "failed" and data["error"]["code"] == "validation_failed"
    assert data["error"]["detail"] == "RuntimeError: disk on fire" and data["progress"]["phase"] == "check"


def test_files_that_cannot_be_validated_are_refused_at_once(client, context):
    response = client.post("/api/validations", json={"job_ids": ["no-such-job"]})
    assert response.status_code == 404
    assert client.get("/api/validations/no-such-session/events").status_code == 404
